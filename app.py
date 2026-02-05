from openpyxl import load_workbook, Workbook
from datetime import datetime
import bcrypt
import os

ARQ_FUNCIONARIOS = "funcionarios.xlsx"
PASTA_PONTOS = "pontos"

# SETUP INICIAL
def inicializar_funcionarios():
    if not os.path.exists(ARQ_FUNCIONARIOS):
        
        planilha = Workbook()
        pagina = planilha.active
        pagina.append(["Nome", "PIN_HASH"])
        planilha.save(ARQ_FUNCIONARIOS)

    if not os.path.exists(PASTA_PONTOS):
        os.mkdir(PASTA_PONTOS)

# FUNCIONÁRIOS
def cadastrar_funcionario(nome, pin):
    planilha = load_workbook(ARQ_FUNCIONARIOS)
    pagina = planilha.active

    # evita duplicar funcionário
    for linha in pagina.iter_rows(min_row=2, values_only=True):
        if linha[0] == nome:
            print(">> Funcionário já cadastrado.")
            return

    pin_hash = bcrypt.hashpw(pin.encode(), bcrypt.gensalt())
    pagina.append([nome, pin_hash.decode()])
    planilha.save(ARQ_FUNCIONARIOS)

    print(">> Funcionário cadastrado com sucesso!")


def validar_pin(nome, pin_digitado):
    planilha = load_workbook(ARQ_FUNCIONARIOS)
    pagina = planilha.active

    for linha in pagina.iter_rows(min_row=2, values_only=True):
        nome_excel, pin_hash = linha
        if nome_excel == nome:
            return bcrypt.checkpw(
                pin_digitado.encode(),
                pin_hash.encode()
            )
    return False

# PONTO
def calcular_horas(entrada, saida):
    formato = "%H:%M"
    entrada_dt = datetime.strptime(entrada, formato)
    saida_dt = datetime.strptime(saida, formato)
    return saida_dt - entrada_dt


def obter_planilha_funcionario(nome):
    caminho = f"{PASTA_PONTOS}/{nome}.xlsx"

    if not os.path.exists(caminho):
        planilha = Workbook()
        pagina = planilha.active
        pagina.append(["Data", "Entrada", "Saída", "Horas"])
        planilha.save(caminho)

    planilha = load_workbook(caminho)
    return planilha, planilha.active, caminho


def registrar_ponto(nome):
    planilha, pagina, caminho = obter_planilha_funcionario(nome)

    agora = datetime.now()
    data_hoje = agora.strftime("%d/%m/%Y")
    hora_atual = agora.strftime("%H:%M")

    ultima_linha_hoje = None

    for linha in pagina.iter_rows(min_row=2):
        data_excel = linha[0].value
        if isinstance(data_excel, datetime):
            data_excel = data_excel.strftime("%d/%m/%Y")

        if data_excel == data_hoje:
            ultima_linha_hoje = linha

    # ENTRADA
    if not ultima_linha_hoje:
        pagina.append([data_hoje, hora_atual, "", ""])
        planilha.save(caminho)
        print(">> Entrada registrada automaticamente!")
        return

    # SAÍDA
    if not ultima_linha_hoje[2].value:
        entrada = ultima_linha_hoje[1].value
        saida = hora_atual

        ultima_linha_hoje[2].value = saida
        horas = calcular_horas(entrada, saida)
        ultima_linha_hoje[3].value = str(horas)

        planilha.save(caminho)
        print(">> Saída registrada automaticamente!")
        return

    print(">> Ponto de hoje já finalizado.")

# MENU
def menu():
    print("\n=== SISTEMA DE PONTO ===")
    print("1 - Registrar ponto")
    print("2 - Cadastrar funcionário")
    print("0 - Sair")


def main():
    inicializar_funcionarios()

    while True:
        menu()
        opcao = input("> Escolha uma opção: ").strip()

        if opcao == "1":
            nome = input("> Nome: ").strip().lower()
            pin = input("> PIN: ").strip()

            if validar_pin(nome, pin):
                registrar_ponto(nome)
            else:
                print(">> Nome ou PIN inválido.")

        elif opcao == "2":
            nome = input("> Novo funcionário: ").strip().lower()
            pin = input("> Crie um PIN: ").strip()
            cadastrar_funcionario(nome, pin)

        elif opcao == "0":
            print(">> Encerrando sistema.")
            break

        else:
            print(">> Opção inválida.")


if __name__ == "__main__":
    main()
